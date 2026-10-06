"""Encrypted document persistence, extraction and evidence boundaries."""
import base64
import io
import json
import os
import zipfile
import httpx
import pytest
from app import main,library,library_extract
from test_backend import client,fake_router
DOCUMENT_KEY="ab"*32

@pytest.fixture
def documents_client(client,monkeypatch):
    monkeypatch.setenv('KENO_DOCUMENT_KEY',DOCUMENT_KEY)
    yield client


def upload(client,text='Confidential orchid budget is 7402.',name='private-orchid.txt',**kwargs):
    raw=text.encode() if isinstance(text,str) else text
    return client.post('/api/v1/library/documents',json={'name':name,'data_base64':base64.b64encode(raw).decode(),'embed':False,**kwargs})


def test_ciphertext_restart_backup_and_original(documents_client):
    from scripts.database import snapshot,validate
    c=documents_client;result=upload(c);assert result.status_code==200,result.text
    doc=result.json()['id'];original=b'Confidential orchid budget is 7402.'
    assert c.get(f'/api/v1/library/documents/{doc}/download').content==original
    snapshot(main.DB_PATH,main.DB_PATH.parent/'backup.db');validate(main.DB_PATH.parent/'backup.db')
    with main.db() as db:
        for table in ('vault_documents','vault_originals'):
            payload=db.execute('SELECT payload FROM '+table).fetchone()[0]
            assert b'orchid' not in payload and b'7402' not in payload
    for file in main.DB_PATH.parent.glob('keno.db*'):assert original not in file.read_bytes()
    os.environ.pop('KENO_DOCUMENT_KEY')
    assert c.get('/api/v1/library/documents').status_code==503
    os.environ['KENO_DOCUMENT_KEY']='cd'*32
    assert c.get('/api/v1/library/documents').status_code==503
    os.environ['KENO_DOCUMENT_KEY']=DOCUMENT_KEY
    assert c.get('/api/v1/library/documents').json()[0]['name']=='private-orchid.txt'
    assert c.post('/api/v1/library/search',json={'question':'orchid budget','mode':'keyword'}).json()['excerpts'][0]['source_id']=='S1'
    assert c.get(f'/api/v1/library/documents/{doc}/download').content==original
    assert c.post('/api/v1/library/unlock',json={'password':'anything'}).status_code==404


def test_reindex_replace_delete_and_tamper(documents_client,monkeypatch):
    async def vectors(texts,kind):return [[1.0]+[0.0]*383 for _ in texts]
    monkeypatch.setattr(library,'embed',vectors)
    first=upload(documents_client).json();assert upload(documents_client).json()['duplicate'] is True
    indexed=upload(documents_client,replace_id=first['id'],embed=True).json()
    assert indexed['version']==2 and indexed['embedding_model']==library.MODEL
    assert documents_client.post('/api/v1/library/search',json={'question':'budget'}).status_code==200
    assert upload(documents_client,'New budget 9400.',replace_id=first['id']).json()['version']==3
    assert b'9400' in documents_client.get('/api/v1/library/documents/'+first['id']+'/download').content
    with main.db() as db:
        data=bytearray(db.execute('SELECT payload FROM vault_documents').fetchone()[0]);data[-1]^=1
        db.execute('UPDATE vault_documents SET payload=?',(bytes(data),))
    assert documents_client.get('/api/v1/library/documents').status_code==422
    assert documents_client.delete('/api/v1/library/documents/'+first['id']).status_code==200
    with main.db() as db:assert db.execute('SELECT count(*) FROM vault_originals').fetchone()[0]==0


def model(answer,calls):
    def handler(request):
        p=json.loads(request.content);calls.append((request.url.path,p))
        if request.url.path=='/apply-template':return httpx.Response(200,json={'prompt':json.dumps(p['messages'])})
        if request.url.path=='/tokenize':return httpx.Response(200,json={'tokens':[1]*100})
        return httpx.Response(200,json={'choices':[{'message':{'content':answer},'finish_reason':'stop'}]})
    return httpx.AsyncClient(base_url='http://llm:8080',transport=httpx.MockTransport(handler))


def test_ask_routing_and_no_plaintext_history(documents_client):
    upload(documents_client);seen=[];calls=[]
    main.app.state.laya=fake_router(seen=seen);main.app.state.llm=model('Budget is 7402. [S1]',calls)
    result=documents_client.post('/api/v1/library/ask',json={'question':'budget','mode':'keyword'})
    assert result.status_code==200,result.text
    assert result.json()['history_saved'] is False and result.json()['citations_present']
    assert len(seen)==1 and set(seen[0]['questions'])=={'thinking','document_scope'}
    assert len([c for c in calls if c[0]=='/v1/chat/completions'])==1
    assert documents_client.get('/api/v1/conversations').json()==[] and documents_client.get('/api/v1/memories').json()==[]
    main.app.state.llm=model('Invalid [S99]',[])
    assert documents_client.post('/api/v1/library/ask',json={'question':'budget','mode':'keyword'}).status_code==502


def test_upload_limits(documents_client):
    assert upload(documents_client,'a'*1_100_000).status_code==422
    assert upload(documents_client,'abc','shortcut.gdoc').status_code==422


def test_office_extractors():
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    book=Workbook();sheet=book.active;sheet.title='Budget';sheet.append(['Orchid',74]);sheet.append(['Formula','=B1*2'])
    stream=io.BytesIO();book.save(stream);r=library_extract.extract('budget.xlsx',stream.getvalue())
    assert r['category']=='spreadsheet' and 'B1: 74' in r['chunks'][0]['text'] and '=B1*2' in r['chunks'][0]['text']
    slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1)).text='Orchid launch';slide.notes_slide.notes_text_frame.text='Speaker notes'
    stream=io.BytesIO();slides.save(stream);r=library_extract.extract('launch.pptx',stream.getvalue())
    assert r['category']=='presentation' and 'Speaker notes' in r['chunks'][0]['text']
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as z:z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Orchid document</w:t></w:r></w:p></w:body></w:document>')
    assert 'Orchid document' in library_extract.extract('notes.docx',stream.getvalue())['chunks'][0]['text']


def test_xml_entity_and_macro_rejection():
    with pytest.raises(Exception):library_extract.extract('bad.xml',b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>')
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as z:z.writestr('word/vbaProject.bin','macro')
    with pytest.raises(Exception):library_extract.extract('bad.docx',stream.getvalue())
